import json
import openpyxl

excel_path = r'e:\softeem\LW\blog\毕设库.xlsx'

wb = openpyxl.load_workbook(excel_path, data_only=True)

result = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n【工作表: {sheet_name}】")
    print(f"行数: {ws.max_row}, 列数: {ws.max_column}")
    
    rows = list(ws.iter_rows(values_only=True))
    
    if not rows:
        print("  (空表)")
        result[sheet_name] = []
        continue
    
    headers = [str(h) if h is not None else f'列{i+1}' for i, h in enumerate(rows[0])]
    print(f"表头: {headers}")
    
    if len(headers) > 5:
        print(f"\n额外字段详情:")
        for i in range(5, len(headers)):
            col_values = [str(row[i]) for row in rows[1:] if i < len(row) and row[i] is not None]
            print(f"  - 【{headers[i]}】: {col_values[:3]} ... (共{len(col_values)}个非空值)")
    
    data = []
    for row in rows[1:]:
        if all(cell is None for cell in row):
            continue
        item = {}
        for i, cell in enumerate(row):
            if i < len(headers):
                item[headers[i]] = str(cell) if cell is not None else ''
        data.append(item)
    
    result[sheet_name] = data
    print(f"\n数据行数: {len(data)}")
    if data:
        print(f"首行完整数据: {json.dumps(data[0], ensure_ascii=False, indent=2)}")

output_path = r'e:\softeem\LW\blog\source\_data\projects_all.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

print(f"\n数据已保存到: {output_path}")
